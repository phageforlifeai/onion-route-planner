"""
export_xlsx.py – Excel route sheets (one tab per vehicle) and the daily order template.
"""
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="7B3F00")   # onion brown
HEAD_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)


def _header(ws, row, headers):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autowidth(ws, min_w=8, max_w=45):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_w, min(max_w, w + 2))


def build_route_workbook(result: dict) -> bytes:
    s = result["summary"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Onion delivery plan – {s['date']}  (departure {s['depart']})"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = s.get("matrix_note", "")
    rows = [
        ("Vehicles used", f"{s['vehicles_used']} of {s['vehicles_total']}"),
        ("Stops served / dropped", f"{s['stops_served']} / {s['stops_dropped']}"),
        ("Total distance (km)", s["total_km"]),
        ("Total driving time (min)", s["total_drive_min"]),
        ("Total load (kg)", f"{s['load_kg']} of {s['fleet_capacity_kg']} ({s['utilisation_pct']}% of fleet)"),
        ("Nearest-neighbour baseline (km)", s.get("baseline_km", "")),
        ("Saving vs baseline", f"{s.get('saving_km', '')} km / {s.get('saving_pct', '')}%"),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v)

    r0 = 4 + len(rows) + 1
    _header(ws, r0, ["Vehicle", "Stops", "Load (kg)", "Capacity", "Util %", "Km", "Drive min",
                     "Service min", "Wait min", "Leave", "Back at depot"])
    for i, rt in enumerate([r for r in result["routes"] if r["used"]], start=1):
        vals = [rt["vehicle"]["name"], len(rt["stops"]), rt["load_kg"], rt["capacity_kg"],
                rt["utilisation_pct"], rt["total_km"], rt["drive_min"], rt["service_min"],
                rt["wait_min"], rt["suggested_departure"], rt["return"]["arrival"]]
        for c, v in enumerate(vals, 1):
            ws.cell(row=r0 + i, column=c, value=v).border = BORDER
    if result.get("dropped"):
        r = r0 + s["vehicles_used"] + 3
        ws.cell(row=r, column=1, value="NOT PLANNED (needs action)").font = Font(bold=True, color="C00000")
        _header(ws, r + 1, ["Customer", "Qty (kg)", "Window", "Reason"])
        for i, d in enumerate(result["dropped"], start=1):
            for c, v in enumerate([d["customer"], d["qty_kg"], d.get("window", ""), d["reason"]], 1):
                ws.cell(row=r + 1 + i, column=c, value=v).border = BORDER
    _autowidth(ws)

    for rt in result["routes"]:
        if not rt["used"]:
            continue
        ws = wb.create_sheet(rt["vehicle"]["name"][:28])
        ws["A1"] = f"{rt['vehicle']['name']} – {s['date']}"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = (f"Leave depot {rt['suggested_departure']}  •  {len(rt['stops'])} stops  •  "
                    f"{rt['load_kg']} kg ({rt['utilisation_pct']}%)  •  {rt['total_km']} km  •  "
                    f"back ≈ {rt['return']['arrival']}")
        _header(ws, 4, ["#", "Trip", "Customer", "Area / Address", "Phone", "Qty (kg)", "ETA", "Window",
                        "Wait (min)", "Leg km", "Cum km", "Remaining on truck (kg)", "Notes", "Delivered ✓"])
        for i, st in enumerate(rt["stops"], start=1):
            vals = [st["seq"], st.get("trip", 1), st["customer"], st.get("address", ""), st.get("phone", ""), st["qty_kg"],
                    st["arrival"], st.get("window", "any"), st["wait_min"], st["leg_km"], st["cum_km"],
                    st["remaining_kg"], st.get("notes", ""), ""]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=4 + i, column=c, value=v)
                cell.border = BORDER
        r = 5 + len(rt["stops"])
        ws.cell(row=r, column=3, value="Final return to depot").font = BOLD
        ws.cell(row=r, column=7, value=rt["return"]["arrival"])
        ws.cell(row=r, column=10, value=rt["return"]["leg_km"])
        ws.cell(row=r, column=11, value=rt["total_km"])
        if len(rt.get("trips", [])) > 1:
            ws.cell(row=r + 1, column=3, value="Trips: " + " | ".join(
                f"Trip {t['trip']}: leave {t['depart']}, {t['load_kg']} kg, {t['total_km']} km, back {t['return']['arrival']}" for t in rt["trips"]))
            r += 1
        ws.cell(row=r + 2, column=1, value="Navigation link(s):").font = BOLD
        for k, u in enumerate(rt["gmaps_urls"]):
            c = ws.cell(row=r + 3 + k, column=1, value=u)
            c.hyperlink = u
            c.font = Font(color="0563C1", underline="single")
        _autowidth(ws)
        ws.column_dimensions["A"].width = 5
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CUSTOMER_HEADERS = ["Customer", "Area", "Address", "Lat", "Lng", "Phone", "Usual kg",
                    "Window From (HH:MM)", "Window To (HH:MM)", "Notes"]
CUSTOMER_KEYS = ["name", "area", "address", "lat", "lng", "phone", "default_qty_kg",
                 "default_tw_from", "default_tw_to", "notes"]


def build_template(customers: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    _header(ws, 1, ["Customer", "Qty (kg)", "Window From (HH:MM)", "Window To (HH:MM)",
                    "Service (min)", "Notes", "Location (address, 'lat, lng' or Google Maps link)"])
    for i, c in enumerate(customers, start=2):
        ws.cell(row=i, column=1, value=c["name"])
        ws.cell(row=i, column=2, value=c.get("default_qty_kg") or None)
        ws.cell(row=i, column=3, value=c.get("default_tw_from") or None)
        ws.cell(row=i, column=4, value=c.get("default_tw_to") or None)
    ws2 = wb.create_sheet("Customers")
    _header(ws2, 1, CUSTOMER_HEADERS)
    for i, c in enumerate(customers, start=2):
        for col, k in enumerate(CUSTOMER_KEYS, 1):
            v = c.get(k, "")
            ws2.cell(row=i, column=col, value=(v if v not in ("", None) else None))
    if customers:
        dv = DataValidation(type="list", formula1=f"=Customers!$A$2:$A${len(customers) + 1}", allow_blank=True)
        dv.error, dv.errorTitle = "Pick a customer from the list (or add a Location for new ones)", "Unknown customer"
        dv.showErrorMessage = False
        ws.add_data_validation(dv)
        dv.add(f"A2:A{max(200, len(customers) + 100)}")
    ws3 = wb.create_sheet("How to use")
    for i, line in enumerate([
        "ORDERS sheet – one row per delivery. Delete rows for customers with no order today.",
        "  Qty in kg. Leave windows blank for 'any time'. Times are 24h, e.g. 06:30.",
        "  Location column: only needed for a NEW customer, or to move an existing one. Best: exact coordinates",
        "  '13.0418, 80.2341' or a Google Maps link (Google Maps → long-press the gate → tap the numbers → copy).",
        "  A written address also works but is less exact (especially without a Google key).",
        "CUSTOMERS sheet – your customer master. Edit anything, add rows for new customers, then upload.",
        "  Lat/Lng = exact drop point and win over the Address text. Blank Lat/Lng + Address = the planner looks it up.",
        "  Names are the key: keep them identical to update a customer; a new name creates a new customer.",
        "Upload this file with 'Import Excel' in the planner – both sheets are read.",
    ], start=1):
        ws3.cell(row=i, column=1, value=line)
    _autowidth(ws)
    _autowidth(ws2)
    ws3.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v):
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%H:%M")
    return str(v).strip()


def parse_customers_xlsx(data: bytes) -> list:
    """Rows of the 'Customers' sheet (if present) as dicts with keys name, area, address, lat, lng, phone,
    default_qty_kg, default_tw_from, default_tw_to, notes – strings, '' when blank."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    if "Customers" not in wb.sheetnames:
        return []
    rows = list(wb["Customers"].iter_rows(values_only=True))
    if not rows:
        return []
    head = [str(h or "").strip().lower() for h in rows[0]]

    def col(*names):
        for n in names:
            for i, h in enumerate(head):
                if h == n or h.startswith(n):
                    return i
        return None

    idx = {"name": col("customer", "name"), "area": col("area"), "address": col("address", "location"),
           "lat": col("lat"), "lng": col("lng", "lon", "long"), "phone": col("phone", "mobile"),
           "default_qty_kg": col("usual", "default qty", "qty", "kg"),
           "default_tw_from": col("window from", "from"), "default_tw_to": col("window to", "to"), "notes": col("note")}
    out = []
    for r in rows[1:]:
        rec = {k: (_cell(r[i]) if i is not None and i < len(r) else "") for k, i in idx.items()}
        if rec["name"]:
            out.append(rec)
    return out


def parse_orders_xlsx(data: bytes) -> list:
    """Return a list of raw order dicts from an uploaded template."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb["Orders"] if "Orders" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [str(h or "").strip().lower() for h in rows[0]]

    def col(*names):
        for n in names:
            for i, h in enumerate(head):
                if h.startswith(n):
                    return i
        return None

    ci, cq, cf, ct, cs, cn, ca = (col("customer", "name"), col("qty", "quantity", "kg"),
                                  col("window from", "from"), col("window to", "to"),
                                  col("service"), col("note"), col("address", "location"))
    out = []
    for r in rows[1:]:
        if ci is None or r[ci] in (None, ""):
            continue

        def g(i):
            if i is None or i >= len(r) or r[i] is None:
                return ""
            v = r[i]
            if hasattr(v, "strftime"):
                return v.strftime("%H:%M")
            return str(v).strip()

        qty = g(cq)
        try:
            qty = float(qty) if qty else 0
        except ValueError:
            qty = 0
        if qty <= 0:
            continue
        svc = g(cs)
        try:
            svc = float(svc) if svc else None
        except ValueError:
            svc = None
        out.append({"customer_name": g(ci), "qty_kg": qty, "tw_from": g(cf), "tw_to": g(ct),
                    "service_min": svc, "notes": g(cn), "address": g(ca)})
    return out
